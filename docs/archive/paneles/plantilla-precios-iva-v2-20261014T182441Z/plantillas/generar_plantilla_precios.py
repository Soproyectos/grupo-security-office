"""Genera la plantilla de revisión del futuro flujo de precios; no importa datos."""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import FormulaRule


TARIFFS = [
    ("INSTALADOR_IVA", "Instalador (con IVA)", "Incluido"),
    ("TIENDA_IVA", "Tienda (con IVA)", "Incluido"),
    ("DPP_ORO_IVA", "DPP Oro (con IVA)", "Incluido"),
    ("DPP_PLATINO_IVA", "DPP Platino (con IVA)", "Incluido"),
    ("CLIENTE_FINAL_IVA", "Cliente Final (con IVA)", "Incluido"),
    ("ORO_SIN_IVA", "Oro (sin IVA)", "No incluido"),
    ("INSTALLER_SIN_IVA", "Installer (sin IVA)", "No incluido"),
]
NAVY = "19364C"
RED = "C90000"
GRAY = "E9EFF3"


def header(ws, labels):
    ws.append(labels)
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 42
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def build(output):
    wb = Workbook()
    prices = wb.active
    prices.title = "PRECIOS"
    header(prices, ["REFERENCIA"] + [code for code, _, _ in TARIFFS])
    prices.column_dimensions["A"].width = 28
    prices["A1"].comment = Comment(
        "Obligatoria. Conservar como texto los ceros iniciales y la referencia exacta. "
        "La referencia debe existir en Biblioteca o quedar pendiente de resolver.",
        "Grupo Security",
    )
    for col, (_, name, tax) in enumerate(TARIFFS, 2):
        letter = prices.cell(1, col).column_letter
        prices.column_dimensions[letter].width = 25
        prices.cell(1, col).comment = Comment(
            f"{name}. IVA: {tax.lower()}. Precio final en COP, no costo. "
            "Celda vacía conserva el precio anterior. No escribir moneda ni fórmulas. "
            "Los decimales recibidos se mostrarán en la revisión antes de redondear.",
            "Grupo Security",
        )
    for row in prices.iter_rows(min_row=2, max_row=2001, min_col=1, max_col=8):
        row[0].number_format = "@"
        for cell in row[1:]:
            # Mostrar decimales de entrada si existen; nunca ocultarlos antes del preview.
            cell.number_format = "#,##0.########"
    prices.add_table(Table(
        displayName="PreciosGrupoSecurity",
        ref="A1:H2001",
        tableStyleInfo=TableStyleInfo(name="TableStyleMedium2", showRowStripes=True),
    ))
    positive = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True)
    positive.errorTitle = "Precio no válido"
    positive.error = "Escribe un número positivo. Vacío conserva el precio; cero no lo elimina."
    positive.promptTitle = "Precio final en COP"
    positive.prompt = "Pega valores, sin fórmulas. El servidor mostrará el redondeo antes de aplicar."
    positive.showErrorMessage = True
    positive.showInputMessage = True
    prices.add_data_validation(positive)
    positive.add("B2:H2001")
    prices.conditional_formatting.add("A2:H2001", FormulaRule(
        formula=['AND($A2="",COUNT($B2:$H2)>0)'],
        fill=PatternFill("solid", fgColor="FDE9E7"),
    ))
    prices.sheet_view.zoomScale = 85
    prices.sheet_properties.pageSetUpPr.fitToPage = True
    prices.page_setup.orientation = "landscape"
    prices.page_setup.paperSize = prices.PAPERSIZE_A3
    prices.page_setup.fitToWidth = 1
    prices.page_setup.fitToHeight = 0
    prices.print_title_rows = "1:1"

    guide = wb.create_sheet("INSTRUCCIONES")
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 110
    guide.merge_cells("B2:D2")
    guide["B2"] = "Grupo Security | Plantilla de precios v1"
    guide["B2"].font = Font(name="Calibri", size=18, bold=True, color=NAVY)
    instructions = [
        ("Estado", "PROPUESTA: el reconocimiento automático y el flujo solo precios todavía no están implementados. No usar esta plantilla como si ya existieran en el importador actual."),
        ("1", "Rellena únicamente PRECIOS. Cada fila necesita REFERENCIA y al menos un precio. Las filas de EJEMPLO son ficticias y no se importan."),
        ("2", "Copia las referencias como texto, conservando ceros, guiones y sufijos. No uses la descripción para identificar un producto."),
        ("3", "Rellena solo las tarifas que actualizas. Dejar una celda vacía conserva el precio existente; no equivale a cero ni a eliminarlo."),
        ("4", "Pega precios finales como valores numéricos de Excel. No pegues fórmulas, costos ni porcentajes. No cambies los encabezados."),
        ("5", "El flujo nuevo propuesto será: subir → revisar cambios y redondeos → aplicar. No deberá pedir mapear esta plantilla."),
        ("COP", "Confirmado: pesos enteros. Propuesta de redondeo: al peso más cercano, ,50 hacia arriba. 123456,49 → 123456; 123456,50 → 123457. Los decimales de entrada se ven antes de aplicar."),
        ("IVA", "Cada columna indica su tratamiento: _IVA ya incluye IVA y _SIN_IVA no lo incluye. No se añadirá IVA ni se calcularán márgenes automáticamente."),
        ("Tarifas", "Esta muestra incluye siete códigos reconocidos por el código local; no confirma que existan o estén activos en la base real. La descarga futura incluirá solo tarifas configuradas y autorizadas."),
        ("Validación", "Referencia desconocida, duplicados, importes no válidos o moneda incompatible se informan sin crear productos ni cambiar sus fichas. La revisión del servidor es obligatoria aunque Excel acepte la celda."),
        ("Capacidad", "La hoja tiene 2000 filas preparadas por comodidad; no fija el límite del futuro importador. Puedes ampliar la tabla conservando su estructura."),
    ]
    for r, (label, text) in enumerate(instructions, 4):
        guide.cell(r, 1, label)
        cell = guide.cell(r, 2, text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = Font(name="Calibri", size=11, color=NAVY)
        guide.row_dimensions[r].height = 54 if len(text) > 185 else 42
        if r % 2 == 0:
            cell.fill = PatternFill("solid", fgColor=GRAY)
    guide.freeze_panes = "B4"

    example = wb.create_sheet("EJEMPLO")
    header(example, ["REFERENCIA FICTICIA", "TARIFA", "PRECIO RECIBIDO", "PRECIO PROPUESTO COP"])
    for values in [
        ("GS-DEMO-001", "INSTALADOR_IVA", 123456.49, 123456),
        ("GS-DEMO-002", "DPP_PLATINO_IVA", 123456.50, 123457),
        ("000123-DEMO", "TIENDA_IVA", 98000, 98000),
    ]:
        example.append(values)
    for letter in "ABCD":
        example.column_dimensions[letter].width = 30
    for row in example.iter_rows(min_row=2, max_row=4):
        row[0].number_format = "@"
        row[2].number_format = "#,##0.00"
        row[3].number_format = "#,##0"
    example["A7"] = "SOLO DEMOSTRACIÓN. No cargar esta hoja ni copiar sus referencias como productos reales."
    example["A7"].font = Font(color=RED, bold=True)

    tariffs = wb.create_sheet("TARIFAS")
    header(tariffs, ["CÓDIGO", "NOMBRE EN CÓDIGO LOCAL", "IVA"])
    for tariff in TARIFFS:
        tariffs.append(tariff)
    tariffs.column_dimensions["A"].width = 30
    tariffs.column_dimensions["B"].width = 38
    tariffs.column_dimensions["C"].width = 22

    contract = wb.create_sheet("CONTRATO")
    header(contract, ["CLAVE", "VALOR"])
    for row in [
        ("template_type", "GS_PRICE_UPDATE"),
        ("schema_version", "1"),
        ("data_sheet", "PRECIOS"),
        ("currency", "COP"),
        ("price_kind", "FINAL_PER_TARIFF"),
        ("decimal_places", "0"),
        ("rounding_proposal", "HALF_UP"),
        ("blank_policy", "KEEP_EXISTING"),
        ("unknown_reference_policy", "PENDING_NO_PRODUCT_CREATION"),
        ("status", "DESIGN_REVIEW_NOT_IMPLEMENTED"),
    ]:
        contract.append(row)
    contract.column_dimensions["A"].width = 34
    contract.column_dimensions["B"].width = 54

    wb.properties.title = "Grupo Security - Actualización de precios - Propuesta v1"
    wb.properties.creator = "Grupo Security"
    wb.properties.description = "Plantilla sintética para revisar el futuro flujo exclusivo de precios."
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)

    # Verificar el artefacto entregado, sin ejecutar importaciones reales.
    checked = load_workbook(output)
    assert checked.sheetnames == ["PRECIOS", "INSTRUCCIONES", "EJEMPLO", "TARIFAS", "CONTRATO"]
    assert [c.value for c in checked["PRECIOS"][1]] == ["REFERENCIA"] + [t[0] for t in TARIFFS]
    assert all(c.value is None for row in checked["PRECIOS"].iter_rows(min_row=2) for c in row)
    assert checked["PRECIOS"]["A2"].number_format == "@"
    assert checked["EJEMPLO"]["A4"].value == "000123-DEMO"
    assert not any(c.data_type == "f" for ws in checked for row in ws for c in row)
    print(f"Plantilla verificada: {output.name}; 7 tarifas de ejemplo; sin datos ni fórmulas en PRECIOS.")


if __name__ == "__main__":
    build(Path(__file__).with_name("PLANTILLA-PRECIOS-GRUPO-SECURITY-v1.xlsx"))
