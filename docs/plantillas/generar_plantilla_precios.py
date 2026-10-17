"""Simplifica la plantilla con IVA conservando las filas editadas por el usuario."""
from pathlib import Path
from io import BytesIO
from hashlib import sha256
from datetime import datetime, timezone
import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation, DataValidationList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

# Encabezados conservados de la edición del usuario. ORO_IVA no es un alias implícito.
TARIFFS = [('INSTALADOR_IVA', 'Instalador'), ('TIENDA_IVA', 'Tienda'),
           ('ORO_IVA', 'Oro'), ('DPP_PLATINO_IVA', 'DPP Platino'),
           ('CLIENTE_FINAL_IVA', 'Cliente final')]
HEADERS = ['REFERENCIA'] + [code for code, _ in TARIFFS]
NAVY, GRAY = '19364C', 'EDF2F5'

def snapshot(ws):
    return {c.coordinate: (c.value, c.data_type) for row in ws.iter_rows(min_row=2)
            for c in row if c.value is not None}

def build(output, source=None):
    source = output if output.exists() else source
    original = source.read_bytes() if source is not None and source.exists() else None
    if original is not None:
        wb = load_workbook(BytesIO(original))
        assert 'PRECIOS' in wb.sheetnames, 'No se encontró PRECIOS; archivo sin modificar.'
        prices = wb['PRECIOS']
        assert [c.value for c in prices[1]] == HEADERS, 'Cambió el encabezado; revisar antes de guardar.'
    else:
        wb = Workbook()
        prices = wb.active
        prices.title = 'PRECIOS'
        prices.append(HEADERS)
    before = snapshot(prices)
    assert not any(kind == 'f' for _, kind in before.values()), 'Revisar fórmulas sin sobrescribirlas.'
    for name in ['INSTRUCCIONES', 'EJEMPLO', 'TARIFAS', 'CONTRATO']:
        if name in wb.sheetnames:
            wb.remove(wb[name])
    assert wb.sheetnames == ['PRECIOS'], 'Hay otras hojas del usuario: revisar antes de quitarlas.'
    last_col, last_row = get_column_letter(len(HEADERS)), max(2001, prices.max_row)
    prices.freeze_panes = 'B2'
    prices.column_dimensions['A'].width = 28
    prices.row_dimensions[1].height = 42
    for c in prices[1]:
        c.fill = PatternFill('solid', fgColor=NAVY)
        c.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        c.alignment = Alignment(vertical='center', wrap_text=True)
    prices['A1'].comment = Comment('Referencia exacta. Conservar como texto ceros y sufijos.', 'Grupo Security')
    for col, (_, name) in enumerate(TARIFFS, 2):
        prices.column_dimensions[get_column_letter(col)].width = 25
        prices.cell(1, col).comment = Comment(
            f'{name}: precio final en COP con IVA incluido. Vacío conserva el precio anterior. '
            'Pegar valores, sin fórmulas. Se revisan los decimales antes de redondear.', 'Grupo Security')
    for row in prices.iter_rows(min_row=2, max_row=last_row, max_col=len(HEADERS)):
        row[0].number_format = '@'
        for c in row[1:]:
            c.number_format = '#,##0.########'  # No ocultar fracciones de entrada.
    prices.tables.clear()
    prices.add_table(Table(displayName='PreciosGrupoSecurity', ref=f'A1:{last_col}{last_row}',
        tableStyleInfo=TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)))
    prices.data_validations = DataValidationList()
    positive = DataValidation(type='decimal', operator='greaterThan', formula1='0', allow_blank=True)
    positive.errorTitle, positive.error = 'Precio no válido', 'Escribe un número positivo. Vacío conserva el precio.'
    positive.promptTitle, positive.prompt = 'Precio final con IVA', 'Pega valores COP; se mostrará el redondeo antes de aplicar.'
    positive.showErrorMessage = positive.showInputMessage = True
    prices.add_data_validation(positive)
    positive.add(f'B2:{last_col}{last_row}')
    prices.conditional_formatting._cf_rules.clear()
    prices.conditional_formatting.add(f'A2:{last_col}{last_row}', FormulaRule(
        formula=[f'AND($A2="",COUNT($B2:${last_col}2)>0)'], fill=PatternFill('solid', fgColor='FDE9E7')))
    prices.sheet_view.zoomScale = 90
    prices.sheet_properties.pageSetUpPr.fitToPage = True
    prices.page_setup.orientation = 'landscape'
    prices.page_setup.paperSize = prices.PAPERSIZE_A3
    prices.page_setup.fitToWidth, prices.page_setup.fitToHeight = 1, 0
    prices.print_title_rows = '1:1'
    guide = wb.create_sheet('INSTRUCCIONES')
    guide.sheet_view.showGridLines = False
    guide.column_dimensions['A'].width, guide.column_dimensions['B'].width = 24, 98
    guide.sheet_view.zoomScale = 95
    headings = {1:'Grupo Security | Cómo preparar tu lista de precios',
                4:'1. Escribe los productos y sus precios',
                10:'2. Dale una última mirada',
                16:'3. Guarda este mismo Excel'}
    for row, title in headings.items():
        guide.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        c = guide.cell(row, 1, title)
        c.font = Font(name='Calibri', size=12, bold=True, color=NAVY)
        c.alignment = Alignment(vertical='center', indent=1)
        guide.row_dimensions[row].height = 29
        if row == 1:
            c.fill = PatternFill('solid', fgColor=NAVY)
            c.font = Font(name='Calibri', size=15, bold=True, color='FFFFFF')
            guide.row_dimensions[row].height = 36
    guide.merge_cells('A2:B2')
    guide['A2'] = 'Trabaja en la hoja PRECIOS. Esta guía te acompaña mientras la llenas.'
    guide['A2'].font, guide['A2'].alignment = Font(name='Calibri', size=11, color=NAVY), Alignment(vertical='center')
    guide.row_dimensions[2].height = 25
    instructions = [
        (5,'Un producto por fila','Copia su referencia completa, con sus letras, guiones y ceros del comienzo. Por ejemplo, 00125 no es lo mismo que 125.'),
        (6,'Precios con IVA','Escribe al menos un precio por producto, según el tipo de cliente. Usa pesos colombianos con IVA incluido; no se sumará otra vez.'),
        (7,'Solo lo que cambia','No tienes que llenar todas las columnas. Deja vacía la casilla del precio que no quieras cambiar: se conservará el que ya tiene el sistema.'),
        (8,'Si copias de otro Excel','Pega el resultado, sin la fórmula. En Excel, elige Pegado especial y luego Valores. Los precios deben ser números, sin el signo $ ni porcentajes.'),
        (11,'Revisa las referencias','No repitas un producto ni cambies los títulos de las columnas. Si la web no encuentra una referencia, te lo indicará para que se revise.'),
        (12,'Revisa los precios','Cada producto debe tener al menos un precio mayor que cero. Si una casilla no se va a actualizar, déjala vacía; no escribas cero.'),
        (13,'Si hay centavos','Los precios quedarán en pesos enteros. Antes de guardarlos en la web podrás revisar cómo quedan después del redondeo.'),
        (14,'Un ejemplo','Con el redondeo propuesto, 123.456,49 quedaría en 123.456 pesos y 123.456,50 quedaría en 123.457 pesos.'),
        (17,'No borres esta guía','Guarda los cambios y conserva las dos hojas: PRECIOS e INSTRUCCIONES. Subirás este mismo archivo de Excel, con el nombre que prefieras.'),
        (18,'Cómo se usará en la web','Abrirás la lista y elegirás Importar. Seleccionarás este Excel, revisarás los cambios y, cuando todo esté bien, podrás aplicar la actualización.')]
    for row, label, text in instructions:
        guide.cell(row, 1, label)
        guide.cell(row, 2, text)
        for c in guide[row]:
            c.font = Font(name='Calibri', size=11, bold=c.column==1, color=NAVY)
            c.alignment = Alignment(wrap_text=True, vertical='center')
            if row == 7:
                c.fill = PatternFill('solid', fgColor='FFF4D6')
            elif row == 14:
                c.fill = PatternFill('solid', fgColor=GRAY)
                c.font = Font(name='Calibri', size=11, italic=True, bold=c.column==1, color=NAVY)
        guide.row_dimensions[row].height = 42 if len(text)>110 else 34
    guide.merge_cells('A20:B20')
    guide['A20'] = 'Esta plantilla está en revisión. La opción para subirla a la web todavía no está habilitada.'
    guide['A20'].font = Font(name='Calibri', size=10, italic=True, color='536471')
    guide['A20'].alignment = Alignment(wrap_text=True, vertical='center')
    guide.row_dimensions[20].height = 34
    guide.freeze_panes, guide.print_area = 'A4', 'A1:B20'
    guide.sheet_properties.pageSetUpPr.fitToPage = True
    guide.page_setup.paperSize = guide.PAPERSIZE_A4
    guide.page_setup.fitToWidth = guide.page_setup.fitToHeight = 1
    wb.active = 0
    wb.properties.title = 'Grupo Security - Precios con IVA - Plantilla v3'
    wb.properties.description = 'Precios finales con IVA. Datos y guía. Formato en revisión.'
    assert snapshot(prices) == before, 'Datos del usuario modificados: se cancela el guardado.'
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(output.stem+'.revision.tmp.xlsx')
    wb.save(temporary)
    checked = load_workbook(temporary)
    assert checked.sheetnames == ['PRECIOS', 'INSTRUCCIONES']
    assert [c.value for c in checked['PRECIOS'][1]] == HEADERS
    assert snapshot(checked['PRECIOS']) == before
    assert checked['PRECIOS']['A2'].number_format == '@'
    checked.close()
    if original is not None:
        assert sha256(source.read_bytes()).digest() == sha256(original).digest(), 'El archivo cambió durante la revisión.'
        archive = output.parent/'archive'
        archive.mkdir(exist_ok=True)
        stamp = datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')
        backup = archive/(source.stem+'-usuario-'+stamp+'.xlsx')
        backup.write_bytes(original)
        print('Respaldo:', backup)
    os.replace(temporary, output)
    print(f'Verificado: 2 hojas, 5 tarifas con IVA y {len(before)} celdas de datos conservadas.')

if __name__ == '__main__':
    build(Path(__file__).with_name('PLANTILLA-PRECIOS-GRUPO-SECURITY-v3.xlsx'),
          source=Path(__file__).with_name('PLANTILLA-PRECIOS-GRUPO-SECURITY-v2.xlsx'))
