#!/usr/bin/env python3
"""Genera el manifiesto auditable de la biblioteca desde una carpeta de Excel.

La carpeta indicada se trata como el universo autorizado: solo se importan los
SKU presentes en esos libros. Las fichas de ``biblioteca/*.json`` enriquecen
los SKU coincidentes, pero nunca agregan productos ajenos a la fuente.

Ejemplo para el lote oficial de 39 listas:

    python scripts/generar_manifiesto_biblioteca.py \
      --lists-dir "D:/Empresa/web-gruposecurity/listas/LISTA_DE_PRECIOS_EXCEL" \
      --expected-files 39

El proceso no copia precios. Produce un manifiesto importable y un reporte de
auditoría con conteos por archivo, clasificación, duplicados y filas omitidas.
"""

from __future__ import annotations

import argparse
import copy
import json
import re
import unicodedata
from collections import Counter, OrderedDict
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_LISTS_DIR = ROOT / "listas" / "LISTA DE PRECIOS EXCEL"
LIBRARY_DIR = ROOT / "biblioteca"
DEFAULT_OUTPUT = ROOT / "biblioteca-import-39-listas.json"
DEFAULT_REPORT = ROOT / "biblioteca-import-39-listas.report.json"
DEFAULT_PUBLIC_MANIFEST = ROOT / "biblioteca-import-web.json"
HEADER_MAX_ROWS = 10
MAX_SKU_LENGTH = 100
ARTIFACT_PREFIXES = ("import-dry-run-", "import-report-", "backup-identidad-")
PRICE_HEADER = re.compile(r"PRECIO|COSTO|VALOR|PREC1|IVA", re.IGNORECASE)


def clean_cell(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def normalize_text(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return "".join(char for char in text if not unicodedata.combining(char)).strip().lower()


def normalize_sku(value: str) -> str:
    return re.sub(r"\s+", "", value).upper()


def workbook_label(file_name: str) -> str:
    label = Path(file_name).stem
    label = re.sub(r"^LISTA\s+", "", label, flags=re.IGNORECASE)
    label = re.sub(r"\s+GRUPO$", "", label, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", label).strip()


def classify_workbook(file_name: str) -> dict[str, str | None]:
    """Clasificación conservadora basada en la lista de procedencia."""

    key = normalize_text(workbook_label(file_name)).upper()

    if "CONTROL" in key and "ACCESO" in key:
        category = "Control de Acceso"
    elif "COMMAX" in key or "INTERCOM" in key:
        category = "Intercomunicación"
    elif "PROTECCION" in key:
        category = "Energía y Respaldo"
    elif "RACK" in key:
        category = "Racks"
    elif any(token in key for token in ("CABLE", "CABLEADO", "LANPRO", "CILES", "DUPLEX")):
        category = "Cableado"
    elif any(token in key for token in ("CONECTIVIDAD", "RUJIE", "UBIQUITI", "TENDA", "TP-LINK")):
        category = "Redes y Cableado"
    elif " IP " in f" {key} ":
        category = "Cámaras IP"
    elif any(token in key for token in ("TURBO", "CCVT", "CCTV", "ALMACENAMIENTO", "DASHCAM", "MONITORES")):
        category = "Videovigilancia"
    elif "ACCESORIOS HIKVISION" in key or "AUDIO Y SENSING" in key:
        category = "Videovigilancia"
    elif any(
        token in key
        for token in (
            "ALARMA",
            "DSC",
            "EBS",
            "SECORLAM",
            "LYNX",
            "EZVIZ",
            "INCENDIO",
            "AUTOMATIZACION",
        )
    ):
        category = "Alarmas y Smart Home"
    else:
        category = "Seguridad electrónica"

    brand: str | None = None
    brand_rules = (
        ("TP-LINK", "TP-Link"),
        ("HIKVISION", "Hikvision"),
        ("HILOOK", "HiLook"),
        ("COMMAX", "Commax"),
        ("INTELBRAS", "Intelbras"),
        ("LANPRO", "LanPro"),
        ("POWEST", "Powest"),
        ("ENERLINE", "Enerline"),
        ("CILES", "Ciles"),
        ("DUPLEX", "Duplex"),
        ("TENDA", "Tenda"),
        ("EZVIZ", "EZVIZ"),
        ("DSC", "DSC"),
        (" ZK", "ZKTeco"),
    )
    is_mixed_brand_list = any(
        token in key
        for token in (
            "VARIOS",
            "RUJIE - UBIQUITI",
            "EBS - SECORLAM - LYNX",
        )
    )
    if not is_mixed_brand_list:
        padded_key = f" {key}"
        for token, candidate in brand_rules:
            if token in padded_key:
                brand = candidate
                break

    return {
        "sourceList": workbook_label(file_name),
        "category": category,
        "brand": brand,
    }


def locate_columns(worksheet: Any) -> tuple[int, int, int, list[int]] | None:
    for row_index, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=HEADER_MAX_ROWS, values_only=True),
        start=1,
    ):
        reference_column = None
        description_column = None
        price_columns: list[int] = []
        for column_index, cell in enumerate(row):
            name = normalize_text(clean_cell(cell))
            if not name:
                continue
            if reference_column is None and re.fullmatch(r"referencia\s*\*?", name):
                reference_column = column_index
            elif description_column is None and re.fullmatch(r"descripcion\s*\*?", name):
                description_column = column_index
            elif PRICE_HEADER.search(name):
                price_columns.append(column_index)
        if reference_column is not None and description_column is not None:
            return row_index, reference_column, description_column, price_columns
    return None


def source_workbooks(lists_dir: Path) -> list[Path]:
    if not lists_dir.is_dir():
        raise FileNotFoundError(f"No existe la carpeta de listas: {lists_dir}")
    return sorted(
        (
            path
            for path in lists_dir.glob("*.xlsx")
            if "PLANTILLA" not in path.name.upper() and not path.name.startswith("~$")
        ),
        key=lambda path: normalize_text(path.name),
    )


def load_source_products(
    lists_dir: Path,
) -> tuple[OrderedDict[str, dict[str, Any]], dict[str, Any]]:
    products: OrderedDict[str, dict[str, Any]] = OrderedDict()
    per_file: list[dict[str, Any]] = []
    omissions: list[dict[str, Any]] = []
    file_sku_sets: dict[str, set[str]] = {}
    totals = Counter()

    for workbook_path in source_workbooks(lists_dir):
        classification = classify_workbook(workbook_path.name)
        workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
        file_skus: set[str] = set()
        file_product_rows = 0
        file_missing_reference = 0
        file_section_rows = 0
        file_long_skus = 0
        sheet_reports: list[dict[str, Any]] = []
        try:
            for worksheet in workbook.worksheets:
                columns = locate_columns(worksheet)
                if not columns:
                    continue
                header_row, reference_column, description_column, price_columns = columns
                sheet_skus: set[str] = set()
                sheet_product_rows = 0
                sheet_missing_reference = 0
                sheet_section_rows = 0
                sheet_long_skus = 0

                for row_number, row in enumerate(
                    worksheet.iter_rows(min_row=header_row + 1, values_only=True),
                    start=header_row + 1,
                ):
                    reference = clean_cell(row[reference_column]) if reference_column < len(row) else None
                    description = (
                        clean_cell(row[description_column])
                        if description_column < len(row)
                        else None
                    )
                    price_values = [
                        clean_cell(row[index]) for index in price_columns if index < len(row)
                    ]
                    prices_empty = all(value is None for value in price_values)

                    if not reference:
                        if description or not prices_empty:
                            sheet_missing_reference += 1
                            file_missing_reference += 1
                            totals["rowsWithoutReferenceOmitted"] += 1
                            omissions.append(
                                {
                                    "file": workbook_path.name,
                                    "sheet": worksheet.title,
                                    "row": row_number,
                                    "reason": "missing_reference",
                                    "description": description,
                                }
                            )
                        continue
                    if not description and prices_empty:
                        sheet_section_rows += 1
                        file_section_rows += 1
                        totals["sectionRowsOmitted"] += 1
                        omissions.append(
                            {
                                "file": workbook_path.name,
                                "sheet": worksheet.title,
                                "row": row_number,
                                "reason": "section_or_non_product_row",
                                "reference": reference,
                            }
                        )
                        continue
                    if len(reference) > MAX_SKU_LENGTH:
                        sheet_long_skus += 1
                        file_long_skus += 1
                        totals["rowsWithLongSkuOmitted"] += 1
                        omissions.append(
                            {
                                "file": workbook_path.name,
                                "sheet": worksheet.title,
                                "row": row_number,
                                "reason": "sku_over_100_characters",
                                "reference": reference,
                                "description": description,
                            }
                        )
                        continue

                    normalized = normalize_sku(reference)
                    if not normalized:
                        continue

                    occurrence = {
                        "reference": reference,
                        "description": description,
                        "file": workbook_path.name,
                        "sheet": worksheet.title,
                        "row": row_number,
                        "sourceList": classification["sourceList"],
                        "category": classification["category"],
                        "brand": classification["brand"],
                    }
                    products.setdefault(normalized, {"occurrences": []})["occurrences"].append(
                        occurrence
                    )
                    file_skus.add(normalized)
                    sheet_skus.add(normalized)
                    file_product_rows += 1
                    sheet_product_rows += 1
                    totals["productRows"] += 1

                sheet_reports.append(
                    {
                        "sheet": worksheet.title,
                        "headerRow": header_row,
                        "productRows": sheet_product_rows,
                        "uniqueSkus": len(sheet_skus),
                        "duplicateRowsWithinSheet": sheet_product_rows - len(sheet_skus),
                        "rowsWithoutReferenceOmitted": sheet_missing_reference,
                        "sectionRowsOmitted": sheet_section_rows,
                        "rowsWithLongSkuOmitted": sheet_long_skus,
                    }
                )
        finally:
            workbook.close()

        file_sku_sets[workbook_path.name] = file_skus
        per_file.append(
            {
                "file": workbook_path.name,
                "sourceList": classification["sourceList"],
                "category": classification["category"],
                "brandHint": classification["brand"],
                "productRows": file_product_rows,
                "uniqueSkus": len(file_skus),
                "duplicateRowsWithinFile": file_product_rows - len(file_skus),
                "rowsWithoutReferenceOmitted": file_missing_reference,
                "sectionRowsOmitted": file_section_rows,
                "rowsWithLongSkuOmitted": file_long_skus,
                "worksheets": sheet_reports,
            }
        )

    files_per_sku = {
        sku: {occurrence["file"] for occurrence in product["occurrences"]}
        for sku, product in products.items()
    }
    for file_report in per_file:
        file_report["duplicateSkusAcrossFiles"] = sum(
            1 for sku in file_sku_sets[file_report["file"]] if len(files_per_sku[sku]) > 1
        )

    totals["workbooks"] = len(per_file)
    totals["worksheetsWithHeaders"] = sum(len(item["worksheets"]) for item in per_file)
    totals["uniqueProducts"] = len(products)
    totals["duplicateOccurrences"] = totals["productRows"] - len(products)
    totals["duplicateSkus"] = sum(
        1 for product in products.values() if len(product["occurrences"]) > 1
    )
    totals["crossFileDuplicateSkus"] = sum(
        1 for files in files_per_sku.values() if len(files) > 1
    )

    return products, {
        "totals": dict(totals),
        "perFile": per_file,
        "omissions": omissions,
    }


def load_curated_library() -> OrderedDict[str, dict[str, Any]]:
    curated: OrderedDict[str, dict[str, Any]] = OrderedDict()
    for file_path in sorted(LIBRARY_DIR.glob("*.json")):
        if file_path.name.startswith(ARTIFACT_PREFIXES):
            continue
        raw = json.loads(file_path.read_text(encoding="utf-8"))
        reference = raw.get("referencia") if isinstance(raw, dict) else None
        if not isinstance(reference, str) or not reference.strip():
            raise ValueError(f"Ficha sin referencia válida: {file_path}")
        normalized = normalize_sku(reference)
        if normalized in curated:
            raise ValueError(f"Referencia curada duplicada: {reference}")
        curated[normalized] = raw
    return curated


def normalized_description(value: str | None) -> str:
    return normalize_text(value)


def select_primary_occurrence(occurrences: list[dict[str, Any]]) -> dict[str, Any]:
    """Elige de forma determinista la fila con la descripción más informativa."""

    return max(
        occurrences,
        key=lambda occurrence: (
            bool(occurrence.get("description")),
            len(occurrence.get("description") or ""),
        ),
    )


def value_block(ficha: dict[str, Any], key: str) -> str | None:
    block = ficha.get(key)
    if not isinstance(block, dict):
        return None
    value = block.get("valor")
    return value.strip() if isinstance(value, str) and value.strip() else None


def build_products(
    source_products: OrderedDict[str, dict[str, Any]],
    curated: OrderedDict[str, dict[str, Any]],
) -> tuple[list[dict[str, Any]], dict[str, Any], list[dict[str, Any]]]:
    products: list[dict[str, Any]] = []
    duplicates: list[dict[str, Any]] = []
    curated_matches = 0
    category_counts: Counter[str] = Counter()
    brand_counts: Counter[str] = Counter()
    state_counts: Counter[str] = Counter()
    source_category_counts: Counter[str] = Counter()
    pending_without_description = 0

    for normalized in sorted(source_products):
        occurrences = source_products[normalized]["occurrences"]
        primary = select_primary_occurrence(occurrences)
        source_labels = list(
            dict.fromkeys(
                f"{occurrence['file']}::{occurrence['sheet']}" for occurrence in occurrences
            )
        )
        source_category_counts[primary["category"]] += 1

        curated_ficha = curated.get(normalized)
        if curated_ficha:
            curated_matches += 1
            ficha = copy.deepcopy(curated_ficha)
            existing_sources = ficha.get("fuentes_consultadas")
            if not isinstance(existing_sources, list):
                existing_sources = []
            ficha["fuentes_consultadas"] = list(
                dict.fromkeys(
                    [source for source in existing_sources if isinstance(source, str)]
                    + source_labels
                )
            )
            if not value_block(ficha, "categoria"):
                ficha["categoria"] = {
                    "valor": primary["category"],
                    "origen": "lista",
                    "razon": f"Clasificación por lista {primary['sourceList']}",
                }
            if primary.get("brand") and not value_block(ficha, "marca"):
                ficha["marca"] = {
                    "valor": primary["brand"],
                    "origen": "lista",
                    "razon": f"Marca única indicada por la lista {primary['sourceList']}",
                }
        else:
            ficha = {
                "referencia": primary["reference"],
                "categoria": {
                    "valor": primary["category"],
                    "origen": "lista",
                    "razon": f"Clasificación por lista {primary['sourceList']}",
                },
                "estado": "pendiente_enriquecimiento",
                "fuentes_consultadas": source_labels,
            }
            if primary.get("brand"):
                ficha["marca"] = {
                    "valor": primary["brand"],
                    "origen": "lista",
                    "razon": f"Marca única indicada por la lista {primary['sourceList']}",
                }
            if primary.get("description"):
                ficha["descripcion"] = {
                    "valor": primary["description"],
                    "origen": "lista",
                }
            else:
                pending_without_description += 1

        products.append(ficha)
        category_counts[value_block(ficha, "categoria") or "Sin categoría"] += 1
        brand_counts[value_block(ficha, "marca") or "Sin marca"] += 1
        state_counts[str(ficha.get("estado") or "sin_estado")] += 1

        if len(occurrences) > 1:
            references = {occurrence["reference"] for occurrence in occurrences}
            descriptions = {
                normalized_description(occurrence.get("description"))
                for occurrence in occurrences
                if occurrence.get("description")
            }
            categories = {occurrence["category"] for occurrence in occurrences}
            brands = {
                occurrence["brand"] for occurrence in occurrences if occurrence.get("brand")
            }
            duplicates.append(
                {
                    "sku": normalized,
                    "occurrenceCount": len(occurrences),
                    "extraOccurrences": len(occurrences) - 1,
                    "hasReferenceVariation": len(references) > 1,
                    "hasDescriptionConflict": len(descriptions) > 1,
                    "hasCategoryConflict": len(categories) > 1,
                    "hasBrandConflict": len(brands) > 1,
                    "selectedSource": {
                        "file": primary["file"],
                        "sheet": primary["sheet"],
                        "row": primary["row"],
                        "reason": "longest_non_empty_description",
                    },
                    "occurrences": occurrences,
                }
            )

    details = {
        "curatedLibraryProducts": len(curated),
        "curatedMatches": curated_matches,
        "curatedOutsideSource": len(curated) - curated_matches,
        "pendingProducts": len(products) - curated_matches,
        "pendingWithoutDescription": pending_without_description,
        "productsWithMultipleOccurrences": len(duplicates),
        "descriptionConflictSkus": sum(
            1 for duplicate in duplicates if duplicate["hasDescriptionConflict"]
        ),
        "categoryConflictSkus": sum(
            1 for duplicate in duplicates if duplicate["hasCategoryConflict"]
        ),
        "brandConflictSkus": sum(1 for duplicate in duplicates if duplicate["hasBrandConflict"]),
        "classification": {
            "effectiveCategories": dict(category_counts.most_common()),
            "effectiveBrands": dict(brand_counts.most_common()),
            "states": dict(state_counts.most_common()),
            "sourceCategories": dict(source_category_counts.most_common()),
        },
    }
    return products, details, duplicates


def reconcile_public_manifest(
    source_products: OrderedDict[str, dict[str, Any]],
    public_manifest_path: Path | None,
) -> dict[str, Any] | None:
    if public_manifest_path is None or not public_manifest_path.is_file():
        return None
    raw = json.loads(public_manifest_path.read_text(encoding="utf-8"))
    records = raw.get("products") if isinstance(raw, dict) else None
    if not isinstance(records, list):
        raise ValueError(f"Snapshot público sin products[]: {public_manifest_path}")

    public_skus: set[str] = set()
    for record in records:
        reference = record.get("referencia") if isinstance(record, dict) else None
        if not isinstance(reference, str) or not reference.strip():
            raise ValueError(f"Snapshot público con ficha sin referencia: {public_manifest_path}")
        normalized = normalize_sku(reference)
        if normalized in public_skus:
            raise ValueError(f"Snapshot público con referencia duplicada: {reference}")
        public_skus.add(normalized)

    official_skus = set(source_products)
    outside = sorted(public_skus - official_skus)
    overlap = public_skus & official_skus
    metadata = raw.get("metadata") if isinstance(raw.get("metadata"), dict) else {}
    return {
        "snapshot": str(public_manifest_path.resolve()),
        "snapshotGeneratedOn": metadata.get("generatedOn"),
        "publicProducts": len(public_skus),
        "presentInOfficialLists": len(overlap),
        "publicOutsideOfficialLists": len(outside),
        "publicOutsideOfficialListSkus": outside,
        "officialListProductsNotPublic": len(official_skus - public_skus),
    }


def build_manifest_and_report(
    lists_dir: Path,
    public_manifest_path: Path | None = None,
) -> tuple[dict[str, Any], dict[str, Any]]:
    source_products, source_audit = load_source_products(lists_dir)
    curated = load_curated_library()
    products, details, duplicates = build_products(source_products, curated)
    public_reconciliation = reconcile_public_manifest(source_products, public_manifest_path)
    source_totals = source_audit["totals"]

    metadata = {
        "generatedOn": date.today().isoformat(),
        "generator": "scripts/generar_manifiesto_biblioteca.py",
        "sourceDirectory": str(lists_dir.resolve()),
        "sourceFiles": source_totals["workbooks"],
        "productRows": source_totals["productRows"],
        "uniqueProducts": source_totals["uniqueProducts"],
        "duplicateOccurrencesMerged": source_totals["duplicateOccurrences"],
        "duplicateSkus": source_totals["duplicateSkus"],
        "rowsWithoutReferenceOmitted": source_totals.get("rowsWithoutReferenceOmitted", 0),
        "sectionRowsOmitted": source_totals.get("sectionRowsOmitted", 0),
        "rowsWithLongSkuOmitted": source_totals.get("rowsWithLongSkuOmitted", 0),
        "curatedProductsMatched": details["curatedMatches"],
        "pendingProducts": details["pendingProducts"],
        "pricesIncluded": False,
    }
    if public_reconciliation:
        metadata.update(
            {
                "publicSnapshotProducts": public_reconciliation["publicProducts"],
                "publicProductsMatched": public_reconciliation["presentInOfficialLists"],
                "publicProductsOutsideOfficialLists": public_reconciliation[
                    "publicOutsideOfficialLists"
                ],
            }
        )
    manifest = {
        "schemaVersion": 1,
        "source": "39 Excel autorizados de LISTA_DE_PRECIOS_EXCEL",
        "metadata": metadata,
        "products": products,
    }
    report = {
        "schemaVersion": 1,
        "generatedOn": date.today().isoformat(),
        "generator": "scripts/generar_manifiesto_biblioteca.py",
        "sourceDirectory": str(lists_dir.resolve()),
        "selectionPolicy": {
            "universe": "Solo SKU presentes en los Excel de la carpeta indicada",
            "duplicateIdentity": "SKU normalizado en mayúsculas y sin espacios",
            "duplicateWinner": "Fila con la descripción no vacía más larga; empate por orden alfabético de archivo",
            "curation": "biblioteca/*.json enriquece solo SKU coincidentes y tiene prioridad",
            "prices": "No se copian precios",
        },
        "summary": {
            **source_totals,
            **{key: value for key, value in details.items() if key != "classification"},
        },
        "classification": details["classification"],
        "perFile": source_audit["perFile"],
        "duplicates": duplicates,
        "omissions": source_audit["omissions"],
    }
    if public_reconciliation:
        report["publicCatalogReconciliation"] = public_reconciliation
    return manifest, report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--lists-dir", type=Path, default=DEFAULT_LISTS_DIR)
    parser.add_argument("--expected-files", type=int)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    parser.add_argument(
        "--public-manifest",
        type=Path,
        default=DEFAULT_PUBLIC_MANIFEST,
        help="Snapshot opcional para cruzar los SKU ya publicados",
    )
    args = parser.parse_args()

    lists_dir = args.lists_dir.resolve()
    workbook_count = len(source_workbooks(lists_dir))
    if args.expected_files is not None and workbook_count != args.expected_files:
        raise ValueError(
            f"Se esperaban {args.expected_files} Excel y se encontraron {workbook_count} en {lists_dir}"
        )

    public_manifest_path = args.public_manifest.resolve() if args.public_manifest else None
    manifest, report = build_manifest_and_report(lists_dir, public_manifest_path)
    output = args.output.resolve()
    report_path = args.report.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    metadata = manifest["metadata"]
    print(f"Manifiesto: {output}")
    print(f"Reporte: {report_path}")
    print(
        f"{metadata['sourceFiles']} archivos | {metadata['productRows']} filas de producto | "
        f"{metadata['uniqueProducts']} SKU únicos | "
        f"{metadata['duplicateOccurrencesMerged']} apariciones duplicadas fusionadas"
    )
    print(
        f"{metadata['curatedProductsMatched']} fichas curadas coincidentes | "
        f"{metadata['pendingProducts']} pendientes de enriquecimiento | "
        f"{metadata['rowsWithoutReferenceOmitted']} filas sin referencia omitidas"
    )


if __name__ == "__main__":
    main()
